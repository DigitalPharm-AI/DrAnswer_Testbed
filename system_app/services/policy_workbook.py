from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from shared.schemas import PolicyWorkbookLoadResult, ResolvedNotificationPolicy, ResolvedPolicyBoundary
from shared.settings import get_settings
from system_app.services.policy_workbook_rows import WorkbookBoundaryRow, WorkbookPolicyRow, WorkbookSystemPolicyRow
from system_app.services.policy_workbook_schema import (
    ALLOWED_PRIMARY_REMINDER_TIMINGS,
    ALLOWED_TEMPLATE_PLACEHOLDERS,
    BOUNDARY_REQUIRED_COLUMNS,
    BOUNDARY_SHEET_NAME,
    DAILY_PATTERN_CONVERSATION_TIME_KEY,
    LEGACY_DAILY_PATTERN_ANALYSIS_TIME_KEY,
    PLACEHOLDER_PATTERN,
    POLICY_COLUMN_DEFAULTS,
    POLICY_REQUIRED_COLUMNS,
    POLICY_SHEET_NAME,
    SYSTEM_POLICY_REQUIRED_COLUMNS,
    SYSTEM_POLICY_SHEET_NAME,
    TEMPLATE_COLUMNS,
    relative_primary_minutes,
)
from system_app.services.policy_workbook_seed import default_boundary_row, default_policy_row, default_system_policy_row

HHMM_PATTERN = re.compile(r"^([01]\d|2[0-3]):[0-5]\d$")
REQUIRED_COLUMNS = POLICY_REQUIRED_COLUMNS
SHEET_NAME = POLICY_SHEET_NAME


class PolicyWorkbookError(ValueError):
    def __init__(self, result: PolicyWorkbookLoadResult) -> None:
        super().__init__("notification policy workbook is invalid")
        self.result = result


class PolicyWorkbookManager:
    def __init__(self, path: Path | None = None) -> None:
        self.path = path or get_settings().policy_workbook_path
        self._policies: list[WorkbookPolicyRow] | None = None
        self._boundaries: list[WorkbookBoundaryRow] | None = None
        self._system_policies: list[WorkbookSystemPolicyRow] | None = None
        self.last_result: PolicyWorkbookLoadResult | None = None

    def reload(self) -> PolicyWorkbookLoadResult:
        self.ensure_workbook_exists()
        self._ensure_workbook_shape()
        policies, policy_errors = self._load_policies()
        boundaries, boundary_errors = self._load_boundaries()
        system_policies, system_policy_errors = self._load_system_policies()
        errors = [*policy_errors, *boundary_errors, *system_policy_errors]
        if not boundary_errors:
            errors.extend(self._validate_policies_against_boundaries(policies, boundaries))
        result = PolicyWorkbookLoadResult(
            path=str(self.path),
            loaded_count=len(policies),
            boundary_loaded_count=len(boundaries),
            system_policy_loaded_count=len(system_policies),
            errors=errors,
        )
        self.last_result = result
        if errors:
            raise PolicyWorkbookError(result)
        self._policies = policies
        self._boundaries = boundaries
        self._system_policies = system_policies
        return result

    def policies(self) -> list[WorkbookPolicyRow]:
        if self._policies is None:
            self.reload()
        return list(self._policies or [])

    def boundaries(self) -> list[WorkbookBoundaryRow]:
        if self._boundaries is None:
            self.reload()
        return list(self._boundaries or [])

    def system_policies(self) -> list[WorkbookSystemPolicyRow]:
        if self._system_policies is None:
            self.reload()
        return list(self._system_policies or [])

    def system_policy_value(self, policy_key: str, default: str) -> str:
        matches = sorted(
            [row for row in self.system_policies() if row.active and row.policy_key == policy_key],
            key=lambda row: row.priority,
            reverse=True,
        )
        return matches[0].value if matches else default

    def resolve_default(self, slot_label: str) -> ResolvedNotificationPolicy:
        active_policies = [policy for policy in self.policies() if policy.active]
        exact = sorted(
            [policy for policy in active_policies if policy.slot_label == slot_label],
            key=lambda policy: policy.priority,
            reverse=True,
        )
        if exact:
            return exact[0].to_resolved_policy(requested_slot_label=slot_label)
        wildcard = sorted(
            [policy for policy in active_policies if policy.slot_label == "*"],
            key=lambda policy: policy.priority,
            reverse=True,
        )
        if wildcard:
            return wildcard[0].to_resolved_policy(requested_slot_label=slot_label)
        result = PolicyWorkbookLoadResult(path=str(self.path), loaded_count=0, errors=["wildcard default policy is missing"])
        raise PolicyWorkbookError(result)

    def resolve_boundary(
        self,
        policy_key: str,
        *,
        requested_slot_label: str,
        fallback_policy_key: str | None = None,
    ) -> ResolvedPolicyBoundary:
        active_boundaries = [boundary for boundary in self.boundaries() if boundary.active]
        boundary = self._resolve_boundary_row(active_boundaries, policy_key)
        if boundary is None and fallback_policy_key:
            boundary = self._resolve_boundary_row(active_boundaries, fallback_policy_key)
        if boundary is None:
            result = PolicyWorkbookLoadResult(
                path=str(self.path),
                loaded_count=len(self._policies or []),
                boundary_loaded_count=0,
                errors=[f'active boundary for policy_key="{policy_key}" is required'],
            )
            raise PolicyWorkbookError(result)
        return boundary.to_resolved_boundary(requested_slot_label=requested_slot_label)

    def max_primary_reminder_lead_minutes(self) -> int:
        active_boundaries = [boundary for boundary in self.boundaries() if boundary.active]
        return max(
            (
                boundary.max_primary_reminder_offset_minutes
                for boundary in active_boundaries
                if "before" in boundary.allowed_primary_reminder_timings
            ),
            default=0,
        )

    def ensure_workbook_exists(self) -> None:
        if self.path.exists():
            return
        self.path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        policy_sheet = workbook.active
        policy_sheet.title = POLICY_SHEET_NAME
        policy_sheet.append(POLICY_REQUIRED_COLUMNS)
        policy_sheet.append(default_policy_row())
        boundary_sheet = workbook.create_sheet(BOUNDARY_SHEET_NAME)
        boundary_sheet.append(BOUNDARY_REQUIRED_COLUMNS)
        boundary_sheet.append(default_boundary_row("default_policy"))
        system_policy_sheet = workbook.create_sheet(SYSTEM_POLICY_SHEET_NAME)
        system_policy_sheet.append(SYSTEM_POLICY_REQUIRED_COLUMNS)
        system_policy_sheet.append(default_system_policy_row())
        workbook.save(self.path)

    def _ensure_workbook_shape(self) -> None:
        workbook = load_workbook(self.path)
        changed = False
        policy_sheet = workbook[POLICY_SHEET_NAME] if POLICY_SHEET_NAME in workbook.sheetnames else None
        if policy_sheet is not None:
            changed = self._append_missing_policy_columns(policy_sheet) or changed
            changed = self._normalize_default_policy_key(policy_sheet) or changed
        if BOUNDARY_SHEET_NAME not in workbook.sheetnames:
            boundary_sheet = workbook.create_sheet(BOUNDARY_SHEET_NAME)
            boundary_sheet.append(BOUNDARY_REQUIRED_COLUMNS)
            for policy_key in self._policy_keys_from_sheet(policy_sheet):
                boundary_sheet.append(default_boundary_row(policy_key))
            changed = True
        else:
            boundary_sheet = workbook[BOUNDARY_SHEET_NAME]
            changed = self._normalize_slot_boundary_sheet(boundary_sheet, policy_sheet) or changed
            changed = self._normalize_default_boundary_key(boundary_sheet) or changed
        if SYSTEM_POLICY_SHEET_NAME not in workbook.sheetnames:
            system_policy_sheet = workbook.create_sheet(SYSTEM_POLICY_SHEET_NAME)
            system_policy_sheet.append(SYSTEM_POLICY_REQUIRED_COLUMNS)
            system_policy_sheet.append(default_system_policy_row())
            changed = True
        else:
            system_policy_sheet = workbook[SYSTEM_POLICY_SHEET_NAME]
            changed = self._append_missing_system_policy_columns(system_policy_sheet) or changed
            changed = self._normalize_system_policy_keys(system_policy_sheet) or changed
            changed = self._ensure_default_system_policy_row(system_policy_sheet) or changed
        if changed:
            try:
                workbook.save(self.path)
            except PermissionError:
                return

    def _append_missing_system_policy_columns(self, sheet) -> bool:
        headers = self._headers(sheet)
        changed = False
        defaults = {
            "policy_key": "",
            "value": "",
            "active": True,
            "priority": 0,
            "description": "",
        }
        for column_name in SYSTEM_POLICY_REQUIRED_COLUMNS:
            if column_name in headers:
                continue
            new_column = sheet.max_column + 1
            sheet.cell(row=1, column=new_column).value = column_name
            for row_number in range(2, sheet.max_row + 1):
                if self._row_is_empty(sheet, row_number):
                    continue
                sheet.cell(row=row_number, column=new_column).value = defaults[column_name]
            changed = True
        return changed

    def _ensure_default_system_policy_row(self, sheet) -> bool:
        headers = self._headers(sheet)
        if "policy_key" not in headers:
            return False
        key_column = headers.index("policy_key") + 1
        for row_number in range(2, sheet.max_row + 1):
            if sheet.cell(row=row_number, column=key_column).value == DAILY_PATTERN_CONVERSATION_TIME_KEY:
                return False
        sheet.append(default_system_policy_row())
        return True

    def _normalize_system_policy_keys(self, sheet) -> bool:
        headers = self._headers(sheet)
        if "policy_key" not in headers:
            return False
        key_column = headers.index("policy_key") + 1
        changed = False
        for row_number in range(2, sheet.max_row + 1):
            if sheet.cell(row=row_number, column=key_column).value == LEGACY_DAILY_PATTERN_ANALYSIS_TIME_KEY:
                sheet.cell(row=row_number, column=key_column).value = DAILY_PATTERN_CONVERSATION_TIME_KEY
                changed = True
        return changed

    def _append_missing_policy_columns(self, sheet) -> bool:
        headers = self._headers(sheet)
        changed = False
        for column_name, default_value in POLICY_COLUMN_DEFAULTS.items():
            if column_name in headers:
                continue
            new_column = sheet.max_column + 1
            sheet.cell(row=1, column=new_column).value = column_name
            for row_number in range(2, sheet.max_row + 1):
                if self._row_is_empty(sheet, row_number):
                    continue
                sheet.cell(row=row_number, column=new_column).value = default_value
            changed = True
        return changed

    def _normalize_default_policy_key(self, sheet) -> bool:
        headers = self._headers(sheet)
        if "policy_key" not in headers:
            return False
        key_column = headers.index("policy_key") + 1
        changed = False
        for row_number in range(2, sheet.max_row + 1):
            if sheet.cell(row=row_number, column=key_column).value == "global_default":
                sheet.cell(row=row_number, column=key_column).value = "default_policy"
                changed = True
        return changed

    def _normalize_slot_boundary_sheet(self, sheet, policy_sheet) -> bool:
        headers = self._headers(sheet)
        if "policy_key" in headers or "slot_label" not in headers:
            return False

        policy_keys = self._policy_keys_from_sheet(policy_sheet)
        policy_keys_by_slot = self._policy_keys_by_slot_from_sheet(policy_sheet)
        source_indexes = {header: headers.index(header) for header in headers if header}
        normalized_rows: list[list[Any]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if all(value in (None, "") for value in row):
                continue
            raw = {header: row[index] if index < len(row) else None for header, index in source_indexes.items()}
            source_slot = str(raw.get("slot_label") or "").strip()
            target_policy_keys = policy_keys if source_slot == "*" else policy_keys_by_slot.get(source_slot, [])
            if not target_policy_keys:
                target_policy_keys = [f"unmatched:{source_slot or raw.get('boundary_key') or 'boundary'}"]
            for policy_key in target_policy_keys:
                normalized_rows.append(
                    [
                        self._normalized_boundary_key(raw.get("boundary_key"), policy_key),
                        policy_key,
                        raw.get("min_extra_reminders"),
                        raw.get("max_extra_reminders"),
                        raw.get("min_interval_minutes"),
                        raw.get("max_interval_minutes"),
                        raw.get("min_missed_dose_after_minutes"),
                        raw.get("max_missed_dose_after_minutes"),
                        raw.get("allowed_primary_reminder_timings"),
                        raw.get("min_primary_reminder_offset_minutes"),
                        raw.get("max_primary_reminder_offset_minutes"),
                        raw.get("active"),
                        raw.get("priority"),
                    ]
                )

        sheet.delete_rows(1, sheet.max_row)
        sheet.append(BOUNDARY_REQUIRED_COLUMNS)
        for row in normalized_rows or [default_boundary_row("default_policy")]:
            sheet.append(row)
        return True

    def _normalize_default_boundary_key(self, sheet) -> bool:
        headers = self._headers(sheet)
        changed = False
        if "policy_key" in headers:
            key_column = headers.index("policy_key") + 1
            for row_number in range(2, sheet.max_row + 1):
                if sheet.cell(row=row_number, column=key_column).value == "global_default":
                    sheet.cell(row=row_number, column=key_column).value = "default_policy"
                    changed = True
        if "boundary_key" in headers:
            boundary_column = headers.index("boundary_key") + 1
            for row_number in range(2, sheet.max_row + 1):
                if sheet.cell(row=row_number, column=boundary_column).value in {"global_boundary", "global_default_boundary"}:
                    sheet.cell(row=row_number, column=boundary_column).value = "default_policy_boundary"
                    changed = True
        return changed

    @staticmethod
    def _normalized_boundary_key(value: Any, policy_key: str) -> str:
        boundary_key = str(value or "").strip()
        if not boundary_key or boundary_key == "global_boundary":
            return f"{policy_key}_boundary"
        return boundary_key

    def _policy_keys_from_sheet(self, sheet) -> list[str]:
        if sheet is None:
            return ["default_policy"]
        rows = self._raw_policy_rows_from_sheet(sheet)
        keys = [
            str(row.get("policy_key")).strip()
            for row in rows
            if self._bool_value(row.get("active")) and str(row.get("policy_key") or "").strip()
        ]
        return list(dict.fromkeys(keys)) or ["default_policy"]

    def _policy_keys_by_slot_from_sheet(self, sheet) -> dict[str, list[str]]:
        if sheet is None:
            return {}
        keys_by_slot: dict[str, list[str]] = {}
        for raw in self._raw_policy_rows_from_sheet(sheet):
            if not self._bool_value(raw.get("active")):
                continue
            slot_label = str(raw.get("slot_label") or "").strip()
            policy_key = str(raw.get("policy_key") or "").strip()
            if not slot_label or not policy_key:
                continue
            keys_by_slot.setdefault(slot_label, []).append(policy_key)
        return {slot: list(dict.fromkeys(keys)) for slot, keys in keys_by_slot.items()}

    def _raw_policy_rows_from_sheet(self, sheet) -> list[dict[str, Any]]:
        headers = self._headers(sheet)
        indexes = {header: headers.index(header) for header in POLICY_REQUIRED_COLUMNS if header in headers}
        rows: list[dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if all(value in (None, "") for value in row):
                continue
            rows.append(
                {
                    column: (
                        row[indexes[column]]
                        if column in indexes and indexes[column] < len(row)
                        else POLICY_COLUMN_DEFAULTS.get(column)
                    )
                    for column in POLICY_REQUIRED_COLUMNS
                }
            )
        return rows

    @staticmethod
    def _headers(sheet) -> list[str]:
        return [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]

    @staticmethod
    def _row_is_empty(sheet, row_number: int) -> bool:
        return all(sheet.cell(row=row_number, column=column).value in (None, "") for column in range(1, sheet.max_column + 1))

    def _load_policies(self) -> tuple[list[WorkbookPolicyRow], list[str]]:
        workbook = load_workbook(self.path)
        if POLICY_SHEET_NAME not in workbook.sheetnames:
            return [], [f"required sheet is missing: {POLICY_SHEET_NAME}"]
        sheet = workbook[POLICY_SHEET_NAME]
        headers = self._headers(sheet)
        missing = [
            column
            for column in POLICY_REQUIRED_COLUMNS
            if column not in headers and column not in POLICY_COLUMN_DEFAULTS
        ]
        if missing:
            return [], [f"missing required columns in {POLICY_SHEET_NAME}: {', '.join(missing)}"]
        column_indexes = {header: headers.index(header) for header in POLICY_REQUIRED_COLUMNS if header in headers}
        policies: list[WorkbookPolicyRow] = []
        errors: list[str] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(value in (None, "") for value in row):
                continue
            raw = {
                column: (
                    row[column_indexes[column]]
                    if column in column_indexes and column_indexes[column] < len(row)
                    else POLICY_COLUMN_DEFAULTS.get(column)
                )
                for column in POLICY_REQUIRED_COLUMNS
            }
            try:
                policies.append(self._parse_policy_row(raw))
            except ValueError as exc:
                errors.append(f"{POLICY_SHEET_NAME} row {row_number}: {exc}")
        if not any(policy.slot_label == "*" and policy.active for policy in policies):
            errors.append('active wildcard policy with slot_label="*" is required')
        return policies, errors

    def _load_boundaries(self) -> tuple[list[WorkbookBoundaryRow], list[str]]:
        workbook = load_workbook(self.path)
        if BOUNDARY_SHEET_NAME not in workbook.sheetnames:
            return [self._parse_boundary_row(dict(zip(BOUNDARY_REQUIRED_COLUMNS, default_boundary_row("default_policy"), strict=True)))], []
        sheet = workbook[BOUNDARY_SHEET_NAME]
        headers = self._headers(sheet)
        missing = [column for column in BOUNDARY_REQUIRED_COLUMNS if column not in headers]
        if missing:
            return [], [f"missing required columns in {BOUNDARY_SHEET_NAME}: {', '.join(missing)}"]
        column_indexes = {header: headers.index(header) for header in BOUNDARY_REQUIRED_COLUMNS}
        boundaries: list[WorkbookBoundaryRow] = []
        errors: list[str] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(value in (None, "") for value in row):
                continue
            raw = {column: row[column_indexes[column]] if column_indexes[column] < len(row) else None for column in BOUNDARY_REQUIRED_COLUMNS}
            try:
                boundaries.append(self._parse_boundary_row(raw))
            except ValueError as exc:
                errors.append(f"{BOUNDARY_SHEET_NAME} row {row_number}: {exc}")
        return boundaries, errors

    def _load_system_policies(self) -> tuple[list[WorkbookSystemPolicyRow], list[str]]:
        workbook = load_workbook(self.path)
        if SYSTEM_POLICY_SHEET_NAME not in workbook.sheetnames:
            return [self._parse_system_policy_row(dict(zip(SYSTEM_POLICY_REQUIRED_COLUMNS, default_system_policy_row(), strict=True)))], []
        sheet = workbook[SYSTEM_POLICY_SHEET_NAME]
        headers = self._headers(sheet)
        missing = [column for column in SYSTEM_POLICY_REQUIRED_COLUMNS if column not in headers]
        if missing:
            return [], [f"missing required columns in {SYSTEM_POLICY_SHEET_NAME}: {', '.join(missing)}"]
        column_indexes = {header: headers.index(header) for header in SYSTEM_POLICY_REQUIRED_COLUMNS}
        rows: list[WorkbookSystemPolicyRow] = []
        errors: list[str] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if all(value in (None, "") for value in row):
                continue
            raw = {column: row[column_indexes[column]] if column_indexes[column] < len(row) else None for column in SYSTEM_POLICY_REQUIRED_COLUMNS}
            try:
                rows.append(self._parse_system_policy_row(raw))
            except ValueError as exc:
                errors.append(f"{SYSTEM_POLICY_SHEET_NAME} row {row_number}: {exc}")
        return rows, errors

    def _parse_policy_row(self, raw: dict[str, Any]) -> WorkbookPolicyRow:
        policy_key = self._required_text(raw.get("policy_key"), "policy_key")
        slot_label = self._required_text(raw.get("slot_label"), "slot_label")
        extra_reminders = self._int_value(raw.get("extra_reminders"), "extra_reminders", minimum=0)
        interval_minutes = self._int_value(raw.get("interval_minutes"), "interval_minutes", minimum=1)
        missed_dose_after_minutes = self._int_value(raw.get("missed_dose_after_minutes"), "missed_dose_after_minutes", minimum=1)
        primary_reminder_timing = self._timing_value(raw.get("primary_reminder_timing"), "primary_reminder_timing")
        primary_reminder_offset_minutes = self._int_value(raw.get("primary_reminder_offset_minutes"), "primary_reminder_offset_minutes", minimum=0)
        if primary_reminder_timing == "at" and primary_reminder_offset_minutes != 0:
            raise ValueError("primary_reminder_offset_minutes must be 0 when primary_reminder_timing is at")
        templates = {column: self._template_value(raw.get(column), column) for column in TEMPLATE_COLUMNS}
        return WorkbookPolicyRow(
            policy_key=policy_key,
            slot_label=slot_label,
            extra_reminders=extra_reminders,
            interval_minutes=interval_minutes,
            missed_dose_after_minutes=missed_dose_after_minutes,
            primary_reminder_timing=primary_reminder_timing,
            primary_reminder_offset_minutes=primary_reminder_offset_minutes,
            medication_title_template=templates["medication_title_template"],
            medication_body_template=templates["medication_body_template"],
            extra_title_template=templates["extra_title_template"],
            extra_body_template=templates["extra_body_template"],
            missed_dose_title_template=templates["missed_dose_title_template"],
            missed_dose_body_template=templates["missed_dose_body_template"],
            active=self._bool_value(raw.get("active")),
            priority=self._int_value(raw.get("priority"), "priority"),
        )

    def _parse_boundary_row(self, raw: dict[str, Any]) -> WorkbookBoundaryRow:
        boundary = WorkbookBoundaryRow(
            boundary_key=self._required_text(raw.get("boundary_key"), "boundary_key"),
            policy_key=self._required_text(raw.get("policy_key"), "policy_key"),
            min_extra_reminders=self._int_value(raw.get("min_extra_reminders"), "min_extra_reminders", minimum=0),
            max_extra_reminders=self._int_value(raw.get("max_extra_reminders"), "max_extra_reminders", minimum=0),
            min_interval_minutes=self._int_value(raw.get("min_interval_minutes"), "min_interval_minutes", minimum=1),
            max_interval_minutes=self._int_value(raw.get("max_interval_minutes"), "max_interval_minutes", minimum=1),
            min_missed_dose_after_minutes=self._int_value(raw.get("min_missed_dose_after_minutes"), "min_missed_dose_after_minutes", minimum=1),
            max_missed_dose_after_minutes=self._int_value(raw.get("max_missed_dose_after_minutes"), "max_missed_dose_after_minutes", minimum=1),
            allowed_primary_reminder_timings=self._timing_list_value(raw.get("allowed_primary_reminder_timings")),
            min_primary_reminder_offset_minutes=self._int_value(
                raw.get("min_primary_reminder_offset_minutes"),
                "min_primary_reminder_offset_minutes",
                minimum=0,
            ),
            max_primary_reminder_offset_minutes=self._int_value(
                raw.get("max_primary_reminder_offset_minutes"),
                "max_primary_reminder_offset_minutes",
                minimum=0,
            ),
            active=self._bool_value(raw.get("active")),
            priority=self._int_value(raw.get("priority"), "priority"),
        )
        self._validate_min_max(boundary.min_extra_reminders, boundary.max_extra_reminders, "extra_reminders")
        self._validate_min_max(boundary.min_interval_minutes, boundary.max_interval_minutes, "interval_minutes")
        self._validate_min_max(boundary.min_missed_dose_after_minutes, boundary.max_missed_dose_after_minutes, "missed_dose_after_minutes")
        self._validate_min_max(
            boundary.min_primary_reminder_offset_minutes,
            boundary.max_primary_reminder_offset_minutes,
            "primary_reminder_offset_minutes",
        )
        return boundary

    def _parse_system_policy_row(self, raw: dict[str, Any]) -> WorkbookSystemPolicyRow:
        policy_key = self._required_text(raw.get("policy_key"), "policy_key")
        if policy_key == DAILY_PATTERN_CONVERSATION_TIME_KEY and (raw.get("value") is None or not str(raw.get("value")).strip()):
            raise ValueError(f"{DAILY_PATTERN_CONVERSATION_TIME_KEY} value is required")
        value = self._required_text(raw.get("value"), "value")
        if policy_key == DAILY_PATTERN_CONVERSATION_TIME_KEY and not HHMM_PATTERN.fullmatch(value):
            raise ValueError(f"{DAILY_PATTERN_CONVERSATION_TIME_KEY} value must use HH:MM format")
        return WorkbookSystemPolicyRow(
            policy_key=policy_key,
            value=value,
            active=self._bool_value(raw.get("active")),
            priority=self._int_value(raw.get("priority"), "priority"),
            description=str(raw.get("description") or "").strip(),
        )

    def _validate_policies_against_boundaries(
        self,
        policies: list[WorkbookPolicyRow],
        boundaries: list[WorkbookBoundaryRow],
    ) -> list[str]:
        active_boundaries = [boundary for boundary in boundaries if boundary.active]
        errors: list[str] = []
        for policy in policies:
            if not policy.active:
                continue
            boundary = self._resolve_boundary_row(active_boundaries, policy.policy_key)
            if boundary is None:
                errors.append(f"{POLICY_SHEET_NAME}:{policy.policy_key} has no policy_key boundary")
                continue
            errors.extend(self._policy_boundary_errors(policy, boundary))
        return errors

    @staticmethod
    def _resolve_boundary_row(boundaries: list[WorkbookBoundaryRow], policy_key: str) -> WorkbookBoundaryRow | None:
        matches = sorted(
            [boundary for boundary in boundaries if boundary.policy_key == policy_key],
            key=lambda boundary: boundary.priority,
            reverse=True,
        )
        return matches[0] if matches else None

    @staticmethod
    def _policy_boundary_errors(policy: WorkbookPolicyRow, boundary: WorkbookBoundaryRow) -> list[str]:
        errors: list[str] = []
        if not boundary.min_extra_reminders <= policy.extra_reminders <= boundary.max_extra_reminders:
            errors.append(f"{POLICY_SHEET_NAME}:{policy.policy_key} extra_reminders is outside policy_boundaries")
        if not boundary.min_interval_minutes <= policy.interval_minutes <= boundary.max_interval_minutes:
            errors.append(f"{POLICY_SHEET_NAME}:{policy.policy_key} interval_minutes is outside policy_boundaries")
        if not boundary.min_missed_dose_after_minutes <= policy.missed_dose_after_minutes <= boundary.max_missed_dose_after_minutes:
            errors.append(f"{POLICY_SHEET_NAME}:{policy.policy_key} missed_dose_after_minutes is outside policy_boundaries")
        if policy.primary_reminder_timing not in boundary.allowed_primary_reminder_timings:
            errors.append(f"{POLICY_SHEET_NAME}:{policy.policy_key} primary_reminder_timing is outside policy_boundaries")
        if not boundary.min_primary_reminder_offset_minutes <= policy.primary_reminder_offset_minutes <= boundary.max_primary_reminder_offset_minutes:
            errors.append(f"{POLICY_SHEET_NAME}:{policy.policy_key} primary_reminder_offset_minutes is outside policy_boundaries")
        if policy.primary_reminder_timing == "at" and policy.primary_reminder_offset_minutes != 0:
            errors.append(f"{POLICY_SHEET_NAME}:{policy.policy_key} at timing must have 0 offset")
        last_alert_minutes = relative_primary_minutes(policy.primary_reminder_timing, policy.primary_reminder_offset_minutes) + (
            policy.extra_reminders * policy.interval_minutes
        )
        if last_alert_minutes > policy.missed_dose_after_minutes:
            errors.append(f"{POLICY_SHEET_NAME}:{policy.policy_key} last medication alert is after missed-dose AI alert")
        return errors

    @staticmethod
    def _validate_min_max(minimum: int, maximum: int, field_name: str) -> None:
        if minimum > maximum:
            raise ValueError(f"{field_name} minimum cannot be greater than maximum")

    @staticmethod
    def _required_text(value: Any, field_name: str) -> str:
        if value is None or not str(value).strip():
            raise ValueError(f"{field_name} is required")
        return str(value).strip()

    @classmethod
    def _template_value(cls, value: Any, field_name: str) -> str:
        text = cls._required_text(value, field_name)
        unknown = sorted(set(PLACEHOLDER_PATTERN.findall(text)) - ALLOWED_TEMPLATE_PLACEHOLDERS)
        if unknown:
            raise ValueError(f"{field_name} has unsupported placeholders: {', '.join(unknown)}")
        return text

    @staticmethod
    def _int_value(value: Any, field_name: str, minimum: int | None = None) -> int:
        try:
            parsed = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{field_name} must be an integer") from exc
        if minimum is not None and parsed < minimum:
            raise ValueError(f"{field_name} must be at least {minimum}")
        return parsed

    @staticmethod
    def _timing_value(value: Any, field_name: str) -> str:
        timing = str(value or "").strip().lower()
        if timing not in ALLOWED_PRIMARY_REMINDER_TIMINGS:
            raise ValueError(f"{field_name} must be one of before, at, after")
        return timing

    @classmethod
    def _timing_list_value(cls, value: Any) -> tuple[str, ...]:
        if value is None or not str(value).strip():
            raise ValueError("allowed_primary_reminder_timings is required")
        import re

        tokens = [token.strip().lower() for token in re.split(r"[,/| ]+", str(value)) if token.strip()]
        unknown = sorted(set(tokens) - ALLOWED_PRIMARY_REMINDER_TIMINGS)
        if unknown:
            raise ValueError(f"allowed_primary_reminder_timings has unsupported values: {', '.join(unknown)}")
        if not tokens:
            raise ValueError("allowed_primary_reminder_timings must include at least one timing")
        return tuple(dict.fromkeys(tokens))

    @staticmethod
    def _bool_value(value: Any) -> bool:
        if isinstance(value, bool):
            return value
        if isinstance(value, int):
            return value != 0
        if isinstance(value, str):
            return value.strip().lower() in {"true", "1", "yes", "y", "active", "활성"}
        return bool(value)


policy_workbook_manager = PolicyWorkbookManager()
