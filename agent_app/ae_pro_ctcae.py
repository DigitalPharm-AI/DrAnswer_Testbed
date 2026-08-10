from __future__ import annotations

import difflib
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import TYPE_CHECKING, Any

import anyio
import openpyxl
from sqlalchemy import Engine, text

from agent_app.embeddings.reference_text import (
    pro_ctcae_embedding_text,
    versioned_pro_ctcae_source,
)
from shared.schemas import AEProCtcaeAssessmentResult, AEProCtcaeQuestion
from shared.settings import get_settings

if TYPE_CHECKING:
    from agent_app.embeddings.base import EmbeddingProvider
    from agent_app.embeddings.semantic_verifier import SemanticMatchVerifier
    from agent_app.persistence.symptom_concept_repository import (
        ClinicalSymptomConceptMatch,
    )

PARSED_ITEMS_SHEET = "Parsed_Items"
OTHER_SYMPTOMS_SHEET = "Other_Symptoms"
KNOWN_SYMPTOM_ALIASES = {
    "Nausea": (
        "메스꺼",
        "속이메스꺼",
        "속울렁",
        "울렁거",
        "구역",
        "오심",
        "속불편",
    ),
    "Dizziness": ("어지럽", "현기증", "핑돌"),
    "Vomiting": ("구토", "토했", "토할"),
    "Diarrhea": ("설사",),
    "Abdominal Pain": ("복통", "배아", "배가아"),
    "Headache": ("두통", "머리아"),
    "Muscle Ache": ("근육통", "근육이아", "쑤심"),
}


@dataclass(frozen=True)
class ProCtcaeQuestionRow:
    symptom_term: str
    korean_symptom_name: str
    item_code: str
    question: str
    response_type: str
    response_options: tuple[str, ...]
    pdf_page: int | None
    sheet_name: str


@dataclass(frozen=True)
class ProCtcaeSymptomEntry:
    symptom_term: str
    korean_symptom_name: str
    sheet_name: str
    questions: tuple[ProCtcaeQuestionRow, ...]

    @property
    def aliases(self) -> tuple[str, ...]:
        values = [self.symptom_term, self.korean_symptom_name]
        values.extend(_parenthetical_aliases(self.korean_symptom_name))
        values.extend(_parenthetical_aliases(self.symptom_term))
        values.extend(KNOWN_SYMPTOM_ALIASES.get(self.symptom_term, ()))
        values.extend(KNOWN_SYMPTOM_ALIASES.get(self.korean_symptom_name, ()))
        return tuple(value for value in values if value)


@dataclass(frozen=True)
class ProCtcaeWorkbook:
    parsed_entries: tuple[ProCtcaeSymptomEntry, ...]
    other_questions: tuple[ProCtcaeQuestionRow, ...]


class ProCtcaeReferenceUnavailable(RuntimeError):
    """Raised when the configured clinical questionnaire cannot be used."""


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _int_or_none(value: object) -> int | None:
    try:
        if value is None or value == "":
            return None
        return int(value)
    except (TypeError, ValueError):
        return None


def _split_options(value: object) -> tuple[str, ...]:
    return tuple(part.strip() for part in _text(value).split("|") if part and part.strip())


def normalize_pro_ctcae_lookup_key(value: str) -> str:
    return re.sub(r"[^0-9a-zA-Z가-힣]+", "", value).lower()


def _normalize(value: str) -> str:
    return normalize_pro_ctcae_lookup_key(value)


def _parenthetical_aliases(value: str) -> list[str]:
    return [match.strip() for match in re.findall(r"\(([^)]+)\)", value or "") if match.strip()]


def _ngrams(value: str, sizes: Iterable[int] = (1, 2, 3)) -> Counter[str]:
    normalized = _normalize(value)
    grams: Counter[str] = Counter()
    for size in sizes:
        if len(normalized) < size:
            continue
        for index in range(0, len(normalized) - size + 1):
            grams[normalized[index : index + size]] += 1
    if not grams and normalized:
        grams[normalized] += 1
    return grams


def _cosine(left: Counter[str], right: Counter[str]) -> float:
    if not left or not right:
        return 0.0
    dot = sum(count * right.get(key, 0) for key, count in left.items())
    left_norm = math.sqrt(sum(count * count for count in left.values()))
    right_norm = math.sqrt(sum(count * count for count in right.values()))
    if left_norm == 0 or right_norm == 0:
        return 0.0
    return dot / (left_norm * right_norm)


def _similarity(left: str, right: str) -> float:
    normalized_left = _normalize(left)
    normalized_right = _normalize(right)
    if not normalized_left or not normalized_right:
        return 0.0
    cosine = _cosine(_ngrams(normalized_left), _ngrams(normalized_right))
    sequence = difflib.SequenceMatcher(None, normalized_left, normalized_right).ratio()
    contains_bonus = 0.08 if normalized_left in normalized_right or normalized_right in normalized_left else 0.0
    return min(1.0, max(cosine, sequence) + contains_bonus)


def _is_exact_match(symptom_text: str, aliases: Iterable[str]) -> bool:
    normalized_input = _normalize(symptom_text)
    if not normalized_input:
        return False
    for alias in aliases:
        normalized_alias = _normalize(alias)
        if not normalized_alias:
            continue
        if normalized_input == normalized_alias:
            return True
        if len(normalized_alias) >= 3 and normalized_alias in normalized_input:
            return True
    return False


def _question_model(row: ProCtcaeQuestionRow) -> AEProCtcaeQuestion:
    return AEProCtcaeQuestion(
        symptom_term=row.symptom_term,
        korean_symptom_name=row.korean_symptom_name,
        item_code=row.item_code,
        question=row.question,
        response_type=row.response_type,
        response_options=list(row.response_options),
        pdf_page=row.pdf_page,
        sheet_name=row.sheet_name,
    )


def _entry_score(symptom_text: str, entry: ProCtcaeSymptomEntry) -> float:
    return max((_similarity(symptom_text, alias) for alias in entry.aliases), default=0.0)


def _candidate(entry: ProCtcaeSymptomEntry, score: float, match_type: str = "similarity") -> dict:
    return {
        "symptom_term": entry.symptom_term,
        "korean_symptom_name": entry.korean_symptom_name,
        "sheet_name": entry.sheet_name,
        "similarity": round(score, 4),
        "match_type": match_type,
    }


def _load_sheet_rows(path: Path, sheet_name: str) -> list[ProCtcaeQuestionRow]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        return []
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [_text(value) for value in rows[0]]
    offset = 1 if "No" in header else 0
    parsed: list[ProCtcaeQuestionRow] = []
    for row in rows[1:]:
        if not row or not any(_text(value) for value in row):
            continue
        parsed.append(
            ProCtcaeQuestionRow(
                symptom_term=_text(row[offset]),
                korean_symptom_name=_text(row[offset + 1]),
                item_code=_normalize(_text(row[offset])) + "_" + _text(row[offset + 2]),
                question=_text(row[offset + 3]),
                response_type=_text(row[offset + 4]),
                response_options=_split_options(row[offset + 5]),
                pdf_page=_int_or_none(row[offset + 6] if len(row) > offset + 6 else None),
                sheet_name=sheet_name,
            )
        )
    return parsed


def _group_entries(rows: Iterable[ProCtcaeQuestionRow]) -> tuple[ProCtcaeSymptomEntry, ...]:
    grouped: dict[tuple[str, str, str], list[ProCtcaeQuestionRow]] = defaultdict(list)
    for row in rows:
        grouped[(row.sheet_name, row.symptom_term, row.korean_symptom_name)].append(row)
    return tuple(
        ProCtcaeSymptomEntry(
            sheet_name=sheet_name,
            symptom_term=symptom_term,
            korean_symptom_name=korean_symptom_name,
            questions=tuple(questions),
        )
        for (sheet_name, symptom_term, korean_symptom_name), questions in grouped.items()
    )


@lru_cache(maxsize=8)
def _load_workbook_cached(path_text: str, modified_time: float) -> ProCtcaeWorkbook:
    del modified_time
    path = Path(path_text)
    parsed_rows = _load_sheet_rows(path, PARSED_ITEMS_SHEET)
    other_rows = _load_sheet_rows(path, OTHER_SYMPTOMS_SHEET)
    return ProCtcaeWorkbook(parsed_entries=_group_entries(parsed_rows), other_questions=tuple(other_rows))


def load_workbook(path: Path | None = None) -> ProCtcaeWorkbook:
    resolved_path = Path(path or get_settings().pro_ctcae_workbook_path)
    if not resolved_path.exists():
        raise ProCtcaeReferenceUnavailable(
            "pro_ctcae_reference_workbook_missing"
        )
    modified_time = resolved_path.stat().st_mtime
    workbook = _load_workbook_cached(str(resolved_path), modified_time)
    if not workbook.parsed_entries:
        raise ProCtcaeReferenceUnavailable(
            "pro_ctcae_reference_workbook_invalid"
        )
    return workbook


def pro_ctcae_source_version(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return versioned_pro_ctcae_source(digest.hexdigest())


def pro_ctcae_equivalent_aliases_for_reference_term(
    term_text: str,
    *,
    workbook_path: Path | None = None,
) -> tuple[str, ...]:
    """Return only aliases belonging to an exact Pro-CTCAE concept.

    This intentionally performs no fuzzy or local-similarity matching. The
    caller must first establish a reference term through exact matching or
    vector retrieval plus LLM verification.
    """

    normalized = normalize_pro_ctcae_lookup_key(term_text)
    if not normalized:
        return ()
    settings = get_settings()
    path = Path(
        workbook_path or settings.pro_ctcae_workbook_path
    ).resolve()
    workbook = load_workbook(path)
    for entry in workbook.parsed_entries:
        if any(
            normalize_pro_ctcae_lookup_key(alias) == normalized
            for alias in entry.aliases
        ):
            return tuple(dict.fromkeys(entry.aliases))
    return ()


async def match_pro_ctcae_symptom_semantic(
    symptom_text: str,
    *,
    engine: Engine,
    embedding_provider: EmbeddingProvider,
    semantic_verifier: SemanticMatchVerifier,
    top_k: int,
    min_similarity: float,
    workbook_path: Path | None = None,
) -> AEProCtcaeAssessmentResult:
    from agent_app.embeddings.semantic_verifier import SemanticCandidate

    settings = get_settings()
    path = Path(
        workbook_path or settings.pro_ctcae_workbook_path
    ).resolve()
    workbook = load_workbook(path)
    clean_symptom = symptom_text.strip()
    if not clean_symptom:
        raise ValueError("pro_ctcae_symptom_text_required")
    source_version = pro_ctcae_source_version(path)
    normalized = normalize_pro_ctcae_lookup_key(clean_symptom)
    exact_rows = await anyio.to_thread.run_sync(
        lambda: _pro_ctcae_exact_rows(
            engine,
            normalized=normalized,
            model_id=embedding_provider.identity.model_id,
            source_version=source_version,
        )
    )
    if exact_rows:
        return _semantic_pro_ctcae_result(
            clean_symptom=clean_symptom,
            workbook=workbook,
            matched_row=exact_rows[0],
            candidate_rows=exact_rows[:top_k],
            match_type="exact",
            scoring_method="exact",
            embedding_provider=embedding_provider.identity.provider,
        )

    query_embedding = await embedding_provider.embed_query(
        pro_ctcae_embedding_text(clean_symptom)
    )
    vector_rows = await anyio.to_thread.run_sync(
        lambda: _pro_ctcae_vector_rows(
            engine,
            embedding=query_embedding,
            model_id=embedding_provider.identity.model_id,
            source_version=source_version,
            top_k=top_k,
            min_similarity=min_similarity,
        )
    )
    selected_ids = await semantic_verifier.matching_candidate_ids(
        clean_symptom,
        [
            SemanticCandidate(
                candidate_id=str(row["embedding_id"]),
                label=(
                    f"{row['alias_text']} | {row['korean_symptom_name']} "
                    f"| {row['symptom_term']}"
                ),
            )
            for row in vector_rows
        ],
        domain="pro_ctcae_symptom_alias",
    )
    matched_row = next(
        (
            row
            for row in vector_rows
            if str(row["embedding_id"]) in selected_ids
        ),
        None,
    )
    if matched_row is not None:
        return _semantic_pro_ctcae_result(
            clean_symptom=clean_symptom,
            workbook=workbook,
            matched_row=matched_row,
            candidate_rows=vector_rows,
            match_type="vector_llm_verified",
            scoring_method="pgvector_cosine_llm",
            embedding_provider=embedding_provider.identity.provider,
        )
    best_similarity = (
        float(vector_rows[0]["similarity"])
        if vector_rows
        else 0.0
    )
    return AEProCtcaeAssessmentResult(
        input_symptom=clean_symptom,
        matched=False,
        match_type="other_symptoms",
        similarity=round(best_similarity, 4),
        threshold=min_similarity,
        scoring_method="pgvector_cosine_llm",
        embedding_provider=embedding_provider.identity.provider,
        sheet_name=OTHER_SYMPTOMS_SHEET,
        questions=[
            _question_model(row) for row in workbook.other_questions
        ],
        candidates=[
            _semantic_candidate(row) for row in vector_rows
        ],
    )


def pro_ctcae_result_for_concept(
    symptom_text: str,
    concept: ClinicalSymptomConceptMatch,
    *,
    workbook_path: Path | None = None,
) -> AEProCtcaeAssessmentResult:
    """Build questionnaire output from an already resolved concept."""

    settings = get_settings()
    path = Path(
        workbook_path or settings.pro_ctcae_workbook_path
    ).resolve()
    workbook = load_workbook(path)
    entry = next(
        (
            candidate
            for candidate in workbook.parsed_entries
            if candidate.symptom_term == concept.symptom_term
            and candidate.korean_symptom_name
            == concept.korean_symptom_name
            and candidate.sheet_name == concept.sheet_name
        ),
        None,
    )
    if entry is None:
        raise ProCtcaeReferenceUnavailable(
            "pro_ctcae_concept_workbook_version_mismatch"
        )
    match_type = (
        "exact"
        if concept.match_type == "EXACT"
        else "vector_llm_verified"
    )
    return AEProCtcaeAssessmentResult(
        input_symptom=symptom_text.strip(),
        matched=True,
        match_type=match_type,
        matched_symptom_term=entry.symptom_term,
        matched_korean_symptom_name=entry.korean_symptom_name,
        similarity=round(concept.similarity, 4),
        threshold=1.0 if match_type == "exact" else 0.0,
        scoring_method="linked_concept_reuse",
        embedding_provider="concept_cache",
        sheet_name=entry.sheet_name,
        questions=[_question_model(row) for row in entry.questions],
        candidates=[
            {
                "symptom_term": entry.symptom_term,
                "korean_symptom_name": entry.korean_symptom_name,
                "sheet_name": entry.sheet_name,
                "similarity": round(concept.similarity, 4),
                "match_type": "linked_concept",
            }
        ],
    )


def _pro_ctcae_exact_rows(
    engine: Engine,
    *,
    normalized: str,
    model_id: str,
    source_version: str,
) -> list[dict[str, Any]]:
    with engine.connect() as connection:
        return list(
            connection.execute(
                text(
                    """
                    SELECT embedding_id, symptom_term,
                           korean_symptom_name, sheet_name,
                           alias_text, 1.0 AS similarity
                    FROM agent_pro_ctcae_alias_embeddings
                    WHERE normalized_alias = :normalized
                      AND model_id = :model_id
                      AND source_version = :source_version
                    ORDER BY embedding_id
                    """
                ),
                {
                    "normalized": normalized,
                    "model_id": model_id,
                    "source_version": source_version,
                },
            ).mappings()
        )


def _pro_ctcae_vector_rows(
    engine: Engine,
    *,
    embedding: list[float],
    model_id: str,
    source_version: str,
    top_k: int,
    min_similarity: float,
) -> list[dict[str, Any]]:
    with engine.connect() as connection:
        connection.execute(
            text("SET LOCAL hnsw.iterative_scan = 'strict_order'")
        )
        raw_rows = list(
            connection.execute(
                text(
                    """
                    SELECT embedding_id, symptom_term,
                           korean_symptom_name, sheet_name, alias_text,
                           1 - (
                               embedding OPERATOR(public.<=>)
                               CAST(:embedding AS public.vector)
                           ) AS similarity
                    FROM agent_pro_ctcae_alias_embeddings
                    WHERE model_id = :model_id
                      AND source_version = :source_version
                    ORDER BY embedding OPERATOR(public.<=>)
                        CAST(:embedding AS public.vector)
                    LIMIT :candidate_limit
                    """
                ),
                {
                    "embedding": json.dumps(
                        embedding,
                        separators=(",", ":"),
                    ),
                    "model_id": model_id,
                    "source_version": source_version,
                    "candidate_limit": min(top_k * 4, 80),
                },
            ).mappings()
        )
    unique_rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for raw_row in raw_rows:
        row = dict(raw_row)
        if float(row["similarity"]) < min_similarity:
            continue
        key = (
            str(row["sheet_name"]),
            str(row["symptom_term"]),
            str(row["korean_symptom_name"]),
        )
        if key in seen:
            continue
        seen.add(key)
        unique_rows.append(row)
        if len(unique_rows) >= top_k:
            break
    return unique_rows


def _semantic_pro_ctcae_result(
    *,
    clean_symptom: str,
    workbook: ProCtcaeWorkbook,
    matched_row: dict[str, Any],
    candidate_rows: list[dict[str, Any]],
    match_type: str,
    scoring_method: str,
    embedding_provider: str,
) -> AEProCtcaeAssessmentResult:
    entry = next(
        (
            candidate
            for candidate in workbook.parsed_entries
            if candidate.symptom_term == str(matched_row["symptom_term"])
            and candidate.korean_symptom_name
            == str(matched_row["korean_symptom_name"])
            and candidate.sheet_name == str(matched_row["sheet_name"])
        ),
        None,
    )
    if entry is None:
        raise ProCtcaeReferenceUnavailable(
            "pro_ctcae_embedding_workbook_version_mismatch"
        )
    similarity = float(matched_row["similarity"])
    return AEProCtcaeAssessmentResult(
        input_symptom=clean_symptom,
        matched=True,
        match_type=match_type,
        matched_symptom_term=entry.symptom_term,
        matched_korean_symptom_name=entry.korean_symptom_name,
        similarity=round(similarity, 4),
        threshold=1.0 if match_type == "exact" else 0.0,
        scoring_method=scoring_method,
        embedding_provider=embedding_provider,
        sheet_name=entry.sheet_name,
        questions=[_question_model(row) for row in entry.questions],
        candidates=[
            _semantic_candidate(row) for row in candidate_rows
        ],
    )


def _semantic_candidate(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "symptom_term": str(row["symptom_term"]),
        "korean_symptom_name": str(row["korean_symptom_name"]),
        "sheet_name": str(row["sheet_name"]),
        "similarity": round(float(row["similarity"]), 4),
        "match_type": (
            "exact"
            if float(row["similarity"]) == 1.0
            else "vector_candidate"
        ),
    }


def match_pro_ctcae_symptom(symptom_text: str, *, threshold: float | None = None, workbook_path: Path | None = None) -> AEProCtcaeAssessmentResult:
    settings = get_settings()
    workbook = load_workbook(workbook_path)
    clean_symptom = symptom_text.strip()
    threshold_value = threshold if threshold is not None else settings.pro_ctcae_similarity_threshold
    scored_entries = sorted(((_entry_score(clean_symptom, entry), entry) for entry in workbook.parsed_entries), key=lambda item: item[0], reverse=True)
    exact_entry = next((entry for _, entry in scored_entries if _is_exact_match(clean_symptom, entry.aliases)), None)
    if exact_entry is not None:
        candidates = [_candidate(exact_entry, 1.0, "exact")]
        candidates.extend(_candidate(entry, score) for score, entry in scored_entries[:4] if entry != exact_entry)
        return AEProCtcaeAssessmentResult(
            input_symptom=clean_symptom,
            matched=True,
            match_type="exact",
            matched_symptom_term=exact_entry.symptom_term,
            matched_korean_symptom_name=exact_entry.korean_symptom_name,
            similarity=1.0,
            threshold=threshold_value,
            scoring_method="exact",
            embedding_provider="",
            sheet_name=exact_entry.sheet_name,
            questions=[_question_model(row) for row in exact_entry.questions],
            candidates=candidates[:5],
        )

    best_score, best_entry = scored_entries[0] if scored_entries else (0.0, None)
    if best_entry is not None and best_score >= threshold_value:
        return AEProCtcaeAssessmentResult(
            input_symptom=clean_symptom,
            matched=True,
            match_type="similarity",
            matched_symptom_term=best_entry.symptom_term,
            matched_korean_symptom_name=best_entry.korean_symptom_name,
            similarity=round(best_score, 4),
            threshold=threshold_value,
            scoring_method="local_similarity",
            embedding_provider="",
            sheet_name=best_entry.sheet_name,
            questions=[_question_model(row) for row in best_entry.questions],
            candidates=[_candidate(entry, score) for score, entry in scored_entries[:5]],
        )

    return AEProCtcaeAssessmentResult(
        input_symptom=clean_symptom,
        matched=False,
        match_type="other_symptoms",
        similarity=round(best_score, 4),
        threshold=threshold_value,
        scoring_method="local_similarity",
        embedding_provider="",
        sheet_name=OTHER_SYMPTOMS_SHEET,
        questions=[_question_model(row) for row in workbook.other_questions],
        candidates=[_candidate(entry, score) for score, entry in scored_entries[:5]],
    )
