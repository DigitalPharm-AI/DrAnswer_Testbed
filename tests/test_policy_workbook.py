from __future__ import annotations

from datetime import date, datetime, time

import pytest
from openpyxl import Workbook, load_workbook

import system_app.services.policy_service as policy_service
from shared.schemas import SystemPolicyDelta
from system_app.models import Notification, ReminderPolicy, SystemPolicyOverride
from system_app.services.dose_event_service import prepare_notification_window
from system_app.services.medication_plan_service import create_medication_plan
from system_app.services.policy_workbook import (
    BOUNDARY_REQUIRED_COLUMNS,
    BOUNDARY_SHEET_NAME,
    REQUIRED_COLUMNS,
    SHEET_NAME,
    SYSTEM_POLICY_REQUIRED_COLUMNS,
    SYSTEM_POLICY_SHEET_NAME,
    PolicyWorkbookError,
    PolicyWorkbookManager,
)
from tests.helpers import build_session


def write_policy_workbook(
    path,
    rows: list[list],
    boundary_rows: list[list] | None = None,
    system_policy_rows: list[list] | None = None,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(REQUIRED_COLUMNS)
    for row in rows:
        sheet.append(row)
    if boundary_rows is not None:
        boundary_sheet = workbook.create_sheet(BOUNDARY_SHEET_NAME)
        boundary_sheet.append(BOUNDARY_REQUIRED_COLUMNS)
        for row in boundary_rows:
            boundary_sheet.append(row)
    if system_policy_rows is not None:
        system_policy_sheet = workbook.create_sheet(SYSTEM_POLICY_SHEET_NAME)
        system_policy_sheet.append(SYSTEM_POLICY_REQUIRED_COLUMNS)
        for row in system_policy_rows:
            system_policy_sheet.append(row)
    workbook.save(path)


def policy_row(
    *,
    key: str = "default_policy",
    slot_label: str = "*",
    extra_reminders: int = 1,
    interval_minutes: int = 30,
    missed_dose_after_minutes: int = 90,
    primary_reminder_timing: str = "at",
    primary_reminder_offset_minutes: int = 0,
    medication_title_template: str = "{medication_name} 복약 알림",
    medication_body_template: str = "{slot_label} {scheduled_time} 복약 시간입니다.",
    extra_title_template: str = "{medication_name} 추가 알림",
    extra_body_template: str = "{slot_label} {scheduled_time} 추가 알림입니다.",
    missed_dose_title_template: str = "미복용 확인",
    missed_dose_body_template: str = "{slot_label} {medication_name} 미복용 확인이 필요합니다.",
    active: bool = True,
    priority: int = 0,
) -> list:
    return [
        key,
        slot_label,
        extra_reminders,
        interval_minutes,
        missed_dose_after_minutes,
        primary_reminder_timing,
        primary_reminder_offset_minutes,
        medication_title_template,
        medication_body_template,
        extra_title_template,
        extra_body_template,
        missed_dose_title_template,
        missed_dose_body_template,
        active,
        priority,
    ]


def boundary_row(
    *,
    key: str = "default_policy_boundary",
    policy_key: str = "default_policy",
    min_extra_reminders: int = 0,
    max_extra_reminders: int = 5,
    min_interval_minutes: int = 1,
    max_interval_minutes: int = 120,
    min_missed_dose_after_minutes: int = 1,
    max_missed_dose_after_minutes: int = 240,
    allowed_primary_reminder_timings: str = "before,at,after",
    min_primary_reminder_offset_minutes: int = 0,
    max_primary_reminder_offset_minutes: int = 120,
    active: bool = True,
    priority: int = 0,
) -> list:
    return [
        key,
        policy_key,
        min_extra_reminders,
        max_extra_reminders,
        min_interval_minutes,
        max_interval_minutes,
        min_missed_dose_after_minutes,
        max_missed_dose_after_minutes,
        allowed_primary_reminder_timings,
        min_primary_reminder_offset_minutes,
        max_primary_reminder_offset_minutes,
        active,
        priority,
    ]


def system_policy_row(
    *,
    key: str = "daily_pattern_conversation_time",
    value: str = "08:30",
    active: bool = True,
    priority: int = 0,
    description: str = "일일 복약 패턴 대화 요청 시각",
) -> list:
    return [key, value, active, priority, description]


def use_policy_workbook(monkeypatch, path) -> PolicyWorkbookManager:
    manager = PolicyWorkbookManager(path)
    monkeypatch.setattr(policy_service, "policy_workbook_manager", manager)
    return manager


def test_policy_workbook_is_seeded_when_missing(tmp_path):
    manager = PolicyWorkbookManager(tmp_path / "default_notification_policies.xlsx")

    result = manager.reload()

    assert result.loaded_count == 1
    assert result.boundary_loaded_count == 1
    assert result.errors == []
    assert manager.path.exists()
    policy = manager.resolve_default("아침 08:00")
    assert policy.source == "excel_default"
    assert policy.slot_label == "아침 08:00"
    assert policy.missed_dose_after_minutes == 90
    assert policy.primary_reminder_timing == "at"
    assert policy.primary_reminder_offset_minutes == 0
    boundary = manager.resolve_boundary(policy.policy_key, requested_slot_label=policy.slot_label)
    assert boundary.max_extra_reminders == 5
    workbook = load_workbook(manager.path)
    assert SYSTEM_POLICY_SHEET_NAME in workbook.sheetnames
    assert manager.system_policy_value("daily_pattern_conversation_time", "00:00") == "08:30"


def test_daily_pattern_conversation_time_uses_system_policy_sheet(tmp_path, monkeypatch):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [policy_row(key="default_policy")],
        boundary_rows=[boundary_row(policy_key="default_policy")],
        system_policy_rows=[system_policy_row(value="09:15")],
    )
    use_policy_workbook(monkeypatch, path)
    policy_service.reload_policy_workbook()

    assert policy_service.daily_pattern_conversation_time() == time(9, 15)


@pytest.mark.parametrize("value", ["8시30분", "25:00", ""])
def test_daily_pattern_conversation_time_rejects_invalid_values(tmp_path, value):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [policy_row(key="default_policy")],
        boundary_rows=[boundary_row(policy_key="default_policy")],
        system_policy_rows=[system_policy_row(value=value)],
    )
    manager = PolicyWorkbookManager(path)

    with pytest.raises(PolicyWorkbookError) as exc_info:
        manager.reload()

    assert "daily_pattern_conversation_time" in " ".join(exc_info.value.result.errors)


def test_daily_pattern_conversation_time_can_use_patient_override(tmp_path, monkeypatch):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [policy_row(key="default_policy")],
        boundary_rows=[boundary_row(policy_key="default_policy")],
        system_policy_rows=[system_policy_row(value="08:30")],
    )
    use_policy_workbook(monkeypatch, path)
    policy_service.reload_policy_workbook()

    with build_session() as session:
        applied, message = policy_service.apply_system_policy_delta(
            session,
            SystemPolicyDelta(
                policy_key="daily_pattern_conversation_time",
                value="09:20",
                reason="환자가 정책 대화 시간을 변경 요청",
                source="patient_request",
            ),
        )

        assert applied is True
        assert message == "일일 패턴 대화 요청 시간을 09:20로 변경했습니다."
        assert policy_service.daily_pattern_conversation_time(session) == time(9, 20)
        view = policy_service.daily_pattern_conversation_time_view(session)
        assert view["policy_key"] == "daily_pattern_conversation_time"
        assert view["value"] == "09:20"
        assert view["source_label"] == "요청 정책"
        assert session.query(SystemPolicyOverride).filter(SystemPolicyOverride.active.is_(True)).count() == 1


def test_policy_workbook_rejects_unknown_template_placeholder(tmp_path):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [
            policy_row(
                medication_body_template="{patient_name}님 복약 시간입니다.",
            )
        ],
    )
    manager = PolicyWorkbookManager(path)

    with pytest.raises(PolicyWorkbookError) as exc_info:
        manager.reload()

    assert "unsupported placeholders" in exc_info.value.result.errors[0]


def test_policy_resolver_prefers_override_then_exact_then_wildcard(tmp_path, monkeypatch):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [
            policy_row(key="wildcard", slot_label="*", extra_reminders=1, interval_minutes=30, missed_dose_after_minutes=90),
            policy_row(key="morning_exact", slot_label="아침 08:00", extra_reminders=2, interval_minutes=20, missed_dose_after_minutes=80, priority=10),
        ],
    )
    use_policy_workbook(monkeypatch, path)
    with build_session() as session:
        exact_policy = policy_service.resolve_policy_for_slot(session, "demo-patient", "아침 08:00", date(2026, 4, 20))
        wildcard_policy = policy_service.resolve_policy_for_slot(session, "demo-patient", "점심 13:00", date(2026, 4, 20))
        session.add(
            ReminderPolicy(
                patient_id="demo-patient",
                policy_key="override",
                slot_label="아침 08:00",
                extra_reminders=3,
                interval_minutes=5,
                missed_dose_after_minutes=25,
                effective_start_date=date(2026, 4, 20),
                effective_end_date=date(2026, 4, 21),
                reason="환자별 override",
                source="patient_request",
                active=True,
            )
        )
        session.commit()
        override_policy = policy_service.resolve_policy_for_slot(session, "demo-patient", "아침 08:00", date(2026, 4, 20))

    assert exact_policy.policy_key == "morning_exact"
    assert exact_policy.extra_reminders == 2
    assert exact_policy.primary_reminder_timing == "at"
    assert wildcard_policy.policy_key == "wildcard"
    assert wildcard_policy.extra_reminders == 1
    assert override_policy.source == "patient_override"
    assert override_policy.policy_key == "override"
    assert override_policy.missed_dose_after_minutes == 25


def test_policy_boundary_resolver_uses_policy_key_not_slot_label(tmp_path):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [
            policy_row(key="wildcard_policy", slot_label="*", extra_reminders=1),
            policy_row(key="morning_policy", slot_label="아침 08:00", extra_reminders=1, priority=10),
        ],
        boundary_rows=[
            boundary_row(key="wildcard_boundary", policy_key="wildcard_policy", max_extra_reminders=3),
            boundary_row(key="morning_boundary", policy_key="morning_policy", max_extra_reminders=1, priority=10),
        ],
    )
    manager = PolicyWorkbookManager(path)
    manager.reload()

    exact_policy = manager.resolve_default("아침 08:00")
    wildcard_policy = manager.resolve_default("점심 13:00")
    exact = manager.resolve_boundary(exact_policy.policy_key, requested_slot_label=exact_policy.slot_label)
    wildcard = manager.resolve_boundary(wildcard_policy.policy_key, requested_slot_label=wildcard_policy.slot_label)

    assert exact.boundary_key == "morning_boundary"
    assert exact.policy_key == "morning_policy"
    assert exact.max_extra_reminders == 1
    assert wildcard.boundary_key == "wildcard_boundary"
    assert wildcard.policy_key == "wildcard_policy"
    assert wildcard.max_extra_reminders == 3


def test_policy_workbook_rejects_policy_without_boundary(tmp_path):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(path, [policy_row(key="default_policy")], boundary_rows=[boundary_row(policy_key="other_policy")])
    manager = PolicyWorkbookManager(path)

    with pytest.raises(PolicyWorkbookError) as exc_info:
        manager.reload()

    assert "notification_policies:default_policy has no policy_key boundary" in exc_info.value.result.errors


def test_policy_workbook_normalizes_slot_boundaries_to_policy_keys(tmp_path):
    path = tmp_path / "default_notification_policies.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(REQUIRED_COLUMNS)
    sheet.append(policy_row(key="wildcard_policy", slot_label="*"))
    sheet.append(policy_row(key="morning_policy", slot_label="아침 08:00", priority=10))
    boundary_sheet = workbook.create_sheet(BOUNDARY_SHEET_NAME)
    slot_boundary_columns = [
        "boundary_key",
        "slot_label",
        "min_extra_reminders",
        "max_extra_reminders",
        "min_interval_minutes",
        "max_interval_minutes",
        "min_missed_dose_after_minutes",
        "max_missed_dose_after_minutes",
        "allowed_primary_reminder_timings",
        "min_primary_reminder_offset_minutes",
        "max_primary_reminder_offset_minutes",
        "active",
        "priority",
    ]
    boundary_sheet.append(slot_boundary_columns)
    boundary_sheet.append(["global_boundary", "*", 0, 4, 1, 120, 1, 240, "before,at,after", 0, 120, True, 0])
    boundary_sheet.append(["morning_slot_boundary", "아침 08:00", 0, 1, 1, 120, 1, 240, "before,at,after", 0, 120, True, 10])
    workbook.save(path)

    manager = PolicyWorkbookManager(path)
    manager.reload()
    normalized = load_workbook(path)
    headers = [cell.value for cell in normalized[BOUNDARY_SHEET_NAME][1]]
    boundary_policy_keys = [row[1] for row in normalized[BOUNDARY_SHEET_NAME].iter_rows(min_row=2, values_only=True)]
    morning = manager.resolve_boundary("morning_policy", requested_slot_label="아침 08:00")

    assert headers == BOUNDARY_REQUIRED_COLUMNS
    assert boundary_policy_keys.count("wildcard_policy") == 1
    assert boundary_policy_keys.count("morning_policy") == 2
    assert morning.boundary_key == "morning_slot_boundary"
    assert morning.max_extra_reminders == 1


def test_policy_workbook_rejects_invalid_boundary_range(tmp_path):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [policy_row()],
        boundary_rows=[boundary_row(min_interval_minutes=30, max_interval_minutes=10)],
    )
    manager = PolicyWorkbookManager(path)

    with pytest.raises(PolicyWorkbookError) as exc_info:
        manager.reload()

    assert "interval_minutes minimum cannot be greater than maximum" in exc_info.value.result.errors[0]


def test_policy_validation_rejects_delta_outside_boundary(tmp_path, monkeypatch):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [policy_row()],
        boundary_rows=[boundary_row(max_extra_reminders=2, max_interval_minutes=30)],
    )
    use_policy_workbook(monkeypatch, path)
    delta = policy_service.NotificationPolicyDelta(
        slot_label="아침 08:00",
        extra_reminders=3,
        interval_minutes=10,
        effective_start_date=date(2026, 4, 20),
        effective_end_date=date(2026, 4, 27),
        reason="테스트",
        source="pattern_analysis",
    )

    is_valid, message = policy_service.validate_policy_delta(delta)

    assert is_valid is False
    assert "추가 알림은 0~2회" in message


def test_notifications_use_resolved_policy_templates_and_missed_delay(tmp_path, monkeypatch):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [
            policy_row(
                extra_reminders=1,
                interval_minutes=15,
                missed_dose_after_minutes=45,
                medication_title_template="{medication_name} 정시",
                medication_body_template="{slot_label} {scheduled_time} 정시 복약입니다.",
                extra_title_template="{medication_name} 한번 더",
                extra_body_template="{slot_label} {scheduled_time} 아직 기록이 없어요.",
                missed_dose_title_template="{medication_name} 확인",
                missed_dose_body_template="{slot_label} {scheduled_time} 미복용 확인이 필요합니다.",
            )
        ],
    )
    use_policy_workbook(monkeypatch, path)
    with build_session() as session:
        create_medication_plan(
            session,
            medication_name="정책약",
            dosage="1정",
            start_date=date(2026, 4, 20),
            end_date=date(2026, 4, 20),
            times_csv="08:00",
            instructions="테스트",
        )

        early_payloads = prepare_notification_window(
            session,
            datetime(2026, 4, 20, 8, 0),
            datetime(2026, 4, 20, 8, 44),
        )
        alerts = session.query(Notification).filter(Notification.notification_type == "medication_alert").order_by(Notification.visible_at).all()
        alert_views = [(alert.visible_at, alert.title, alert.body) for alert in alerts]
        missed_payloads = prepare_notification_window(
            session,
            datetime(2026, 4, 20, 8, 44),
            datetime(2026, 4, 20, 8, 45),
        )
        missed_alert = session.query(Notification).filter(Notification.notification_type == "conversation_alert").one()
        missed_alert_view = (missed_alert.title, missed_alert.body)
        missed_detected_at = missed_payloads[0].detected_at if missed_payloads else None

    assert early_payloads == []
    assert [alert[0] for alert in alert_views] == [datetime(2026, 4, 20, 8, 0), datetime(2026, 4, 20, 8, 15)]
    assert alert_views[0][1] == "정책약 정시"
    assert alert_views[0][2] == "아침 08:00 08:00 정시 복약입니다."
    assert alert_views[1][1] == "정책약 한번 더"
    assert len(missed_payloads) == 1
    assert missed_detected_at == datetime(2026, 4, 20, 8, 45)
    assert missed_alert_view[0] == "AI가 대화를 요청합니다."
    assert missed_alert_view[1] == "아침 08:00 08:00 미복용 확인이 필요합니다."


def test_primary_reminder_can_be_scheduled_before_dose_time(tmp_path, monkeypatch):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [
            policy_row(
                extra_reminders=1,
                interval_minutes=10,
                missed_dose_after_minutes=90,
                primary_reminder_timing="before",
                primary_reminder_offset_minutes=30,
            )
        ],
        boundary_rows=[boundary_row()],
    )
    use_policy_workbook(monkeypatch, path)
    with build_session() as session:
        create_medication_plan(
            session,
            medication_name="선행알림약",
            dosage="1정",
            start_date=date(2026, 4, 20),
            end_date=date(2026, 4, 20),
            times_csv="08:00",
            instructions="테스트",
        )

        prepare_notification_window(
            session,
            datetime(2026, 4, 20, 7, 29),
            datetime(2026, 4, 20, 7, 30),
        )
        alerts = session.query(Notification).filter(Notification.notification_type == "medication_alert").order_by(Notification.visible_at).all()

    assert [alert.visible_at for alert in alerts] == [datetime(2026, 4, 20, 7, 30), datetime(2026, 4, 20, 7, 40)]


def test_primary_reminder_can_be_scheduled_after_dose_time(tmp_path, monkeypatch):
    path = tmp_path / "default_notification_policies.xlsx"
    write_policy_workbook(
        path,
        [
            policy_row(
                extra_reminders=1,
                interval_minutes=10,
                missed_dose_after_minutes=90,
                primary_reminder_timing="after",
                primary_reminder_offset_minutes=30,
            )
        ],
        boundary_rows=[boundary_row()],
    )
    use_policy_workbook(monkeypatch, path)
    with build_session() as session:
        create_medication_plan(
            session,
            medication_name="후행알림약",
            dosage="1정",
            start_date=date(2026, 4, 20),
            end_date=date(2026, 4, 20),
            times_csv="08:00",
            instructions="테스트",
        )

        prepare_notification_window(
            session,
            datetime(2026, 4, 20, 8, 0),
            datetime(2026, 4, 20, 8, 29),
        )
        early_alerts = session.query(Notification).filter(Notification.notification_type == "medication_alert").all()
        prepare_notification_window(
            session,
            datetime(2026, 4, 20, 8, 29),
            datetime(2026, 4, 20, 8, 30),
        )
        alerts = session.query(Notification).filter(Notification.notification_type == "medication_alert").order_by(Notification.visible_at).all()

    assert early_alerts == []
    assert [alert.visible_at for alert in alerts] == [datetime(2026, 4, 20, 8, 30), datetime(2026, 4, 20, 8, 40)]
